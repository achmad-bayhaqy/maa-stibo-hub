#!/usr/bin/env python3
"""Deep-dive 'Master Data Dictionary (MAA).xlsx'."""
import openpyxl, os
from collections import Counter

PATH = "/home/z/my-project/upload/Master Data Dictionary (MAA).xlsx"
wb = openpyxl.load_workbook(PATH, data_only=True, read_only=False)

def dump(sheet, r1=1, r2=None, maxc=None, trunc=90):
    ws = wb[sheet]
    r2 = r2 or ws.max_row
    maxc = maxc or ws.max_column
    for row in ws.iter_rows(min_row=r1, max_row=r2, max_col=maxc):
        parts = [f"{c.coordinate}={str(c.value)[:trunc]}" for c in row
                 if c.value is not None and str(c.value).strip()]
        if parts: print(f"r{row[0].row:>4}: " + " | ".join(parts))

print("="*100); print("### INSTRUCTIONS (full)"); print("="*100)
dump("Instructions")

print("\n" + "="*100); print("### VERSION HISTORY (full)"); print("="*100)
dump("Version History")

print("\n" + "="*100); print("### VALIDATION BASE TYPE DEFINITION (full)"); print("="*100)
dump("Validation Base Type Definition")

print("\n" + "="*100); print("### SIMPLE LOVs (full)"); print("="*100)
dump("Simple LOVs")

print("\n" + "="*100); print("### BUSINESS RULES (full)"); print("="*100)
dump("Business Rules")

ws = wb["Core Attributes"]
print("\n" + "="*100); print(f"### CORE ATTRIBUTES ({ws.max_row}x{ws.max_column})"); print("="*100)
hdr = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
print("HEADER row1:")
for i, h in enumerate(hdr, 1):
    print(f"  col{i:>2}: {h!r}")
# find header row (maybe row1 or row2)
for hr in (1, 2):
    if ws.cell(row=hr, column=1).value and "attribute" in str(ws.cell(row=hr, column=1).value).lower():
        break

n_data = 0
rows = []
for row in ws.iter_rows(min_row=2, max_col=ws.max_column):
    if any(c.value is not None and str(c.value).strip() for c in row):
        n_data += 1
        rows.append(row)
print(f"\nData rows (non-empty): {n_data}")
print("\nSAMPLE rows 2,3,10,50,120,200,300,420:")
for idx in [2, 3, 10, 50, 120, 200, 300, min(420, len(rows)+1)]:
    row = rows[idx-2]
    print(f"--- data row #{idx} (sheet row {row[0].row}) ---")
    for c in row:
        if c.value is not None and str(c.value).strip():
            print(f"    {hdr[c.column-1]!r}: {str(c.value)[:100]}")

# column fill stats + distinct counts for key columns
print("\nColumn fill counts:")
for c in range(1, ws.max_column + 1):
    vals = [ws.cell(row=r, column=c).value for r in range(2, ws.max_row + 1)]
    nn = [v for v in vals if v is not None and str(v).strip()]
    if nn:
        print(f"  col{c:>2} {str(hdr[c-1])[:44]:46s} non-empty={len(nn):>4} distinct={len(set(map(str,nn)))}")

# counts: validation types, mandatory, LOV refs
def colidx(name_part):
    for i, h in enumerate(hdr):
        if h and name_part.lower() in str(h).lower(): return i + 1
    return None
vbt = colidx("Validation"); mand = colidx("Mandatory"); lovr = colidx("LOV")
print("\nUsing columns -> Validation:", vbt, "Mandatory:", mand, "LOV:", lovr)
if vbt:
    cc = Counter(str(ws.cell(row=r, column=vbt).value) for r in range(2, ws.max_row+1)
                 if ws.cell(row=r, column=vbt).value)
    print("Validation base type distribution:", dict(cc.most_common()))
if mand:
    cc = Counter(str(ws.cell(row=r, column=mand).value) for r in range(2, ws.max_row+1)
                 if ws.cell(row=r, column=mand).value not in (None, ""))
    print("Mandatory flag distribution:", dict(cc.most_common(12)))
if lovr:
    cc = Counter("yes" if ws.cell(row=r, column=lovr).value else "no" for r in range(2, ws.max_row+1))
    print(f"LOV ref col populated: {dict(cc)}")

# LOV inventory
print("\n" + "="*100); print("### LOV SHEET INVENTORY"); print("="*100)
tot = 0
lov_sheets = []
for name in wb.sheetnames:
    ws2 = wb[name]
    # heuristics: not core/metadata sheets
    if name in ("Instructions", "Version History", "Core Attributes", "Validation Base Type Definition",
                "Simple LOVs", "Business Rules", "Core Attributes (Old)", "Article LOVs (Not in Use)"):
        continue
    is_lov = "LOV" in name or name.startswith("SIS") or name in ("Brand Group", "Franchise LOV", "Golf Club Flex")
    entries = 0
    for row in ws2.iter_rows(min_row=2, max_col=2):
        if any(c.value is not None and str(c.value).strip() for c in row): entries += 1
    lov_sheets.append((name, ws2.max_row, ws2.max_column, entries))
    tot += entries
for name, mr, mc, e in lov_sheets:
    print(f"  {name!r:38s} rows={mr:>6} cols={mc:>2} entries={e:>6}")
print(f"Total candidate LOV/reference sheets: {len(lov_sheets)}; total data entries: {tot}")

print("\n" + "="*100); print("### LOV SAMPLES"); print("="*100)
for sheet, r2 in [("Company Code LOV", 12), ("SBU LOV", 10), ("Brand LOV", 12), ("Gender LOV", 9),
                  ("Season LOV", 9), ("Color Code LOV", 15), ("Size Code LOV", 12), ("Material LOV", 10),
                  ("Vendor LOV", 8), ("UOM LOV", 8), ("Retail Price Currency LOV", 6), ("SAP Product Flag LOV", 6)]:
    print(f"\n--- {sheet} ---")
    dump(sheet, 1, r2, trunc=60)
