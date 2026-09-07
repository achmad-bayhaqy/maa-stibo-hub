#!/usr/bin/env python3
"""Deep-dive 'Brand Input files & Naming convention.xlsx' (both sheets, full)."""
import openpyxl

PATH = "/home/z/my-project/upload/Brand Input files & Naming convention.xlsx"
wb = openpyxl.load_workbook(PATH, data_only=True)

for name in wb.sheetnames:
    ws = wb[name]
    print("=" * 100)
    print(f"### SHEET {name!r}  ({ws.max_row} x {ws.max_column})")
    print("=" * 100)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        parts = [f"{c.coordinate}={str(c.value)[:60]}" for c in row
                 if c.value is not None and str(c.value).strip()]
        if parts: print(f"r{row[0].row:>3}: " + " | ".join(parts))
