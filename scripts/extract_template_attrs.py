#!/usr/bin/env python3
"""Extract full attribute list (row 3+) from Template sheet — these are the columns the portal must show."""
import openpyxl, json

path = "/home/z/my-project/upload/NEW - Brand mapping files Template.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb["Template"]

attrs = []
for r in range(3, ws.max_row + 1):
    name = ws.cell(row=r, column=1).value   # Stibo Attribute
    aid = ws.cell(row=r, column=2).value    # Stibo Attribute ID
    val = ws.cell(row=r, column=3).value    # Stibo Validation
    cluster = ws.cell(row=r, column=4).value
    group = ws.cell(row=r, column=5).value
    if name is None and aid is None:
        continue
    attrs.append({
        "row": r,
        "name": str(name).strip() if name else "",
        "id": str(aid).strip() if aid else "",
        "validation": str(val).strip() if val else "",
        "cluster": str(cluster).strip() if cluster else "",
        "grouping": str(group).strip() if group else "",
    })

print(f"TOTAL ATTRIBUTES: {len(attrs)}")
for a in attrs:
    print(f"  r{a['row']:>3} | {a['name']:<38} | {a['id']:<32} | {a['validation']:<12} | {a['cluster']}")

with open("/home/z/my-project/scripts/template_attrs.json", "w") as f:
    json.dump(attrs, f, indent=2)
print("saved -> scripts/template_attrs.json")
