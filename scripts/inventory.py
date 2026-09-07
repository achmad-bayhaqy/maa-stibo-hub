#!/usr/bin/env python3
"""Inventory all 4 workbooks: sheet names + dimensions (read_only, fast)."""
import openpyxl, os, sys

UP = "/home/z/my-project/upload"
FILES = [
    "NEW - Brand mapping files Template.xlsx",
    "Master Data Dictionary (MAA).xlsx",
    "Brand Input files & Naming convention.xlsx",
    "0888-SP-ELLESSE-Recap Sample MAA Sport Licensed-Multi-SP2027-ID-1.xlsx",
]

for f in FILES:
    p = os.path.join(UP, f)
    print("=" * 100)
    print(f"FILE: {f}  ({os.path.getsize(p)/1e6:.1f} MB)")
    print("=" * 100)
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    for ws in wb.worksheets:
        print(f"  {ws.title!r:60s} max_row={ws.max_row:>6} max_col={ws.max_column:>4} state={ws.sheet_state}")
    wb.close()
