#!/usr/bin/env python3
"""Deep recon: MDD sheet list, Ellesse sheet dump, naming sheet, mapping types census"""
import os
from openpyxl import load_workbook
from collections import Counter

UP = "/home/z/my-project/upload"

print("========== 1. MDD workbook: all sheet names ==========")
wb = load_workbook(os.path.join(UP, "Master Data Dictionary (MAA).xlsx"), read_only=True, data_only=True)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"  {sn!r}  {ws.max_row}x{ws.max_column}")
wb.close()

print("\n========== 2. Template workbook: Field Mapping Type census ==========")
wb = load_workbook(os.path.join(UP, "NEW - Brand mapping files Template.xlsx"), read_only=True, data_only=True)
census = Counter()
for sn in wb.sheetnames:
    ws = wb[sn]
    for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 300), min_col=7, max_col=7, values_only=True):
        v = row[0]
        if v: census[str(v).strip()[:60]] += 1
for k, c in census.most_common(30):
    print(f"  {c:5d}  {k!r}")

print("\n========== 3. Ellesse(Licensed) full dump (cols 1-12) ==========")
ws = wb["Ellesse(Licensed)"]
rows = list(ws.iter_rows(min_row=1, max_row=60, max_col=12, values_only=True))
for i, row in enumerate(rows, 1):
    vals = [str(v)[:40].replace("\n", " ") if v is not None else "" for v in row]
    if any(vals):
        print(f"  r{i}: " + " | ".join(vals))

print("\n========== 4. Row 60-236 non-empty source fields (Ellesse) ==========")
count = 0
for i, row in enumerate(ws.iter_rows(min_row=60, max_row=236, max_col=12, values_only=True), 60):
    attr = row[1]; src = row[9]
    if attr:
        count += 1
        if count <= 40:
            print(f"  r{i}: {str(row[0])[:22]} | {str(attr)[:22]} | {str(row[2])[:6]} | type={str(row[6])[:24]} | src={str(src)[:28]}")
print(f"  TOTAL attribute rows: {count}")
wb.close()

print("\n========== 5. Naming convention 'Stibo' sheet full (cols 1-8, rows 1-89) ==========")
wb2 = load_workbook(os.path.join(UP, "Brand Input files & Naming convention.xlsx"), read_only=True, data_only=True)
ws = wb2["Stibo"]
for i, row in enumerate(ws.iter_rows(min_row=1, max_row=89, max_col=8, values_only=True), 1):
    vals = [str(v)[:45].replace("\n", " ") if v is not None else "" for v in row]
    if any(vals):
        print(f"  r{i}: " + " || ".join(v for v in vals if v))
wb2.close()
