#!/usr/bin/env python3
"""Deep-dive sample input file 0888-SP-ELLESSE-... .xlsx (read_only mode)."""
import openpyxl, os

PATH = "/home/z/my-project/upload/0888-SP-ELLESSE-Recap Sample MAA Sport Licensed-Multi-SP2027-ID-1.xlsx"
wb = openpyxl.load_workbook(PATH, read_only=True, data_only=True)

def dump(sheet, r1=1, r2=None, maxc=None, trunc=55):
    ws = wb[sheet]
    r2 = r2 or min(ws.max_row, 100)
    maxc = maxc or ws.max_column
    print(f"--- {sheet!r} dims {ws.max_row}x{ws.max_column}, dumping rows {r1}..{r2} cols<= {maxc} ---")
    n = 0
    for i, row in enumerate(ws.iter_rows(min_row=r1, max_row=r2, max_col=maxc), start=r1):
        parts = [f"col{c.column}={str(c.value)[:trunc]}" for c in row
                 if c.value is not None and str(c.value).strip()]
        if parts:
            n += 1
            print(f"r{i:>3}: " + " | ".join(parts))
    return n

print("SHEETS:", wb.sheetnames)
n = dump("ELL", 1, 4)
print(f"({n} non-empty rows)\n")
n = dump("Mapping", 1, 89)
print(f"({n} non-empty rows)\n")
n = dump("MAPPING GEN ART", 1, 50)
print(f"({n} non-empty rows)\n")
n = dump("Sheet2", 1, 29)
print(f"({n} non-empty rows)\n")
n = dump("Sheet1", 1, 12, maxc=31)
print(f"({n} non-empty rows)\n")
n = dump("Database Stibo", 1, 12)
print(f"({n} non-empty rows)\n")
n = dump("Sheet5", 1, 30)
n = dump("Sheet3", 1, 42)
n = dump("Sheet4", 1, 2)
wb.close()
